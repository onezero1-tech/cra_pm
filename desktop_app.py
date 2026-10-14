# -*- coding: utf-8 -*-
"""
Excel Matrix Pro - Windows 独立免安装桌面应用程序
基于 Tkinter & 原生 Python 数据处理引擎，可通过 PyInstaller 编译为单文件 .exe
"""

import io
import os
import re
import sys
import glob
import shutil
import zipfile
import tempfile
import threading
from copy import copy
from typing import List, Dict, Optional, Tuple
from collections import OrderedDict
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# 依赖数据处理库
import pandas as pd
import openpyxl
from openpyxl import load_workbook
import xlsxwriter
try:
    import py7zr
    _HAS_PY7ZR = True
except ImportError:
    _HAS_PY7ZR = False


# ============================================================
#                      核心业务逻辑算法
# ============================================================

def decode_zip_filename(raw: str) -> str:
    try:
        raw_bytes = raw.encode('cp437')
    except UnicodeEncodeError:
        return raw
    for encoding in ['gbk', 'utf-8', 'cp437']:
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw


def extract_archive_to_dir(archive_path: str, target_dir: str) -> List[str]:
    lower = archive_path.lower()
    extracted: List[str] = []

    def safe_path(base_dir: str, name: str) -> Optional[str]:
        name = name.replace('\\', '/')
        parts = [p for p in name.split('/') if p and p not in ('.', '..')]
        return os.path.join(base_dir, *parts) if parts else None

    if lower.endswith('.zip'):
        with zipfile.ZipFile(archive_path, 'r') as z:
            for info in z.infolist():
                if info.is_dir():
                    continue
                name = decode_zip_filename(info.filename)
                if not name.lower().endswith(('.xls', '.xlsx', '.xlsm')):
                    continue
                out = safe_path(target_dir, name)
                if not out:
                    continue
                os.makedirs(os.path.dirname(out), exist_ok=True)
                with open(out, 'wb') as f:
                    f.write(z.read(info))
                extracted.append(out)
    elif lower.endswith('.7z') and _HAS_PY7ZR:
        with py7zr.SevenZipFile(archive_path, mode='r') as z:
            for fname, bio in z.readall().items():
                if not fname.lower().endswith(('.xls', '.xlsx', '.xlsm')):
                    continue
                out = safe_path(target_dir, fname)
                if not out:
                    continue
                os.makedirs(os.path.dirname(out), exist_ok=True)
                with open(out, 'wb') as f:
                    f.write(bio.read())
                extracted.append(out)
    elif lower.endswith(('.xls', '.xlsx', '.xlsm')):
        dest = os.path.join(target_dir, os.path.basename(archive_path))
        shutil.copy2(archive_path, dest)
        extracted.append(dest)
    else:
        raise ValueError("只支持 .zip / .7z 压缩包或 .xlsx Excel 文件")

    if not extracted:
        raise ValueError("压缩包内未找到 Excel 表格文件")
    return extracted


def norm(s):
    return '' if s is None else re.sub(r'\s+', '', str(s).strip())


def canon_key(s):
    if s is None:
        return ''
    s = str(s).replace('*', '-').replace('（', '(').replace('）', ')').replace('：', ':')
    return re.sub(r'\s+', '', s.strip())


def text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return str(v).strip()


def to_num(v):
    if v is None or v == '' or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def strip_paren(s):
    return re.sub(r'[（(][^）)]*[）)]\s*$', '', str(s)).strip()


def copy_style_no_fill(src_cell, dst_cell):
    dst_cell.font = copy(src_cell.font)
    dst_cell.border = copy(src_cell.border)
    dst_cell.alignment = copy(src_cell.alignment)
    dst_cell.number_format = copy(src_cell.number_format)
    dst_cell.protection = copy(src_cell.protection)


# ---------- 1. PM 拆分算法 ----------

def execute_pm_split(data_path, tpl_path, out_zip_path, sheet_name="02-项目汇总表",
                     usecols="4,5,6,9,11", header_row=1, data_start=4, log_fn=print):
    positions = [int(p.strip()) - 1 for p in usecols.split(',') if p.strip().isdigit()]
    if len(positions) < 2:
        raise ValueError("提取列索引至少需要 2 列")

    log_fn(f"[信息] 正在读取数据文件: {os.path.basename(data_path)} 工作表: {sheet_name}...")
    df_all = pd.read_excel(data_path, sheet_name=sheet_name, header=header_row - 1)
    df = df_all.iloc[:, positions].copy()
    if df.empty:
        raise ValueError("未从指定列中提取到有效数据")

    group_col = df.columns[-1]
    first_col = df.columns[0]

    with open(tpl_path, 'rb') as f:
        tpl_content = f.read()

    with zipfile.ZipFile(out_zip_path, 'w', zipfile.ZIP_DEFLATED) as zip_out:
        groups = list(df.groupby(group_col, sort=False))
        total_groups = len(groups)
        log_fn(f"[信息] 识别到 {total_groups} 个分组...")

        for idx, (k_value, sub_df) in enumerate(groups, 1):
            safe_name = str(k_value).replace('/', '_').replace('\\', '_')
            log_fn(f"[{idx}/{total_groups}] 处理分组: {safe_name} ({len(sub_df)} 行)...")

            tpl_io = io.BytesIO(tpl_content)
            wb = load_workbook(tpl_io)
            ws_tpl = wb['A'] if 'A' in wb.sheetnames else wb[wb.sheetnames[0]]

            for d_value, mini_df in sub_df.groupby(first_col, sort=False):
                sheet_name_d = str(d_value)[:31]
                if sheet_name_d in wb.sheetnames:
                    wb.remove(wb[sheet_name_d])
                ws = wb.copy_worksheet(ws_tpl)
                ws.title = sheet_name_d

                for r_idx, (_, row) in enumerate(mini_df.iterrows(), start=data_start):
                    for c_idx in range(1, min(4, len(row))):
                        val = row.iloc[c_idx]
                        if pd.notna(val):
                            c = ws.cell(row=r_idx, column=c_idx, value=val)
                            copy_style_no_fill(ws_tpl.cell(row=data_start, column=c_idx), c)

            if ws_tpl.title in wb.sheetnames:
                try:
                    wb.remove(ws_tpl)
                except Exception:
                    pass

            out_io = io.BytesIO()
            wb.save(out_io)
            zip_out.writestr(f"{safe_name}.xlsx", out_io.getvalue())

    log_fn(f"[成功] PM 拆分完成，已输出至: {out_zip_path}")


# ---------- 2. PM 合并算法 ----------

def execute_pm_merge(archive_path, output_excel_path, log_fn=print):
    log_fn(f"[信息] 正在解压并读取合并文件: {os.path.basename(archive_path)}...")
    with tempfile.TemporaryDirectory() as tmpdir:
        files = extract_archive_to_dir(archive_path, tmpdir)
        all_frames = []

        for idx, fpath in enumerate(files, 1):
            base = os.path.basename(fpath)
            parts = base.split('-')
            file_key = parts[4] if len(parts) > 4 else base
            if file_key.lower().endswith('.xlsx'):
                file_key = file_key[:-5]

            try:
                xl = pd.ExcelFile(fpath)
                hs_sheets = [s for s in xl.sheet_names if str(s).upper().startswith("HS")]
                for sheet in hs_sheets:
                    df = pd.read_excel(xl, sheet_name=sheet, header=None)
                    if df.shape[1] <= 1:
                        continue
                    df = df.dropna(axis=1, how="all").iloc[3:, 1:]
                    df = df[df.iloc[:, 0].notna()]
                    df.insert(0, "Sheet", sheet)
                    df.insert(0, "File", file_key)
                    all_frames.append(df)
            except Exception as e:
                log_fn(f"[警告] 读取文件失败 {base}: {e}")

        if not all_frames:
            raise ValueError("未找到任何符合条件的 HS 工作表有效数据")

        final = pd.concat(all_frames, ignore_index=True)
        final.to_excel(output_excel_path, index=False)
        log_fn(f"[成功] PM 合并完成！共生成 {len(final)} 行记录 -> {output_excel_path}")


# ---------- 3. CRA 横表合并算法 ----------

def execute_cra_horizontal(archive_path, output_excel_path, template_path=None, log_fn=print):
    log_fn("[信息] 开始执行 CRA 横表宽表合并...")
    with tempfile.TemporaryDirectory() as tmpdir:
        files = extract_archive_to_dir(archive_path, tmpdir)
        data_files = [f for f in files if f.lower().endswith(('.xlsx', '.xlsm', '.xls')) and not os.path.basename(f).startswith('~$')]
        if not data_files:
            raise ValueError("未找到待合并的 Excel 文件")

        # 模板选择
        tpl_file = template_path if (template_path and os.path.isfile(template_path)) else None
        if not tpl_file:
            for f in data_files:
                if '模板' in os.path.basename(f) or 'template' in os.path.basename(f).lower():
                    tpl_file = f
                    break
        if not tpl_file:
            tpl_file = data_files[0]

        log_fn(f"[信息] 使用模板参考源: {os.path.basename(tpl_file)}")

        # 简单提取 CRA 核心结构
        all_rows = []
        for f in data_files:
            try:
                wb = openpyxl.load_workbook(f, data_only=True)
                if 'CRA填写表单' not in wb.sheetnames:
                    continue
                ws = wb['CRA填写表单']
                row1 = [ws.cell(1, c).value for c in range(1, 20)]
                info = {'源文件': os.path.basename(f)}
                for i in range(0, min(16, len(row1)), 2):
                    if row1[i]:
                        info[str(row1[i])] = str(row1[i+1] or '')

                # 读取工作项统计
                for r in range(3, ws.max_row + 1):
                    work_name = str(ws.cell(r, 2).value or '').strip()
                    count_val = to_num(ws.cell(r, 8).value)
                    if work_name and count_val and count_val > 0:
                        info[f"{work_name}_count"] = count_val
                        info[f"{work_name}_备注"] = str(ws.cell(r, 6).value or '')
                all_rows.append(info)
            except Exception as e:
                log_fn(f"[警告] 处理文件失败 {os.path.basename(f)}: {e}")

        if not all_rows:
            raise ValueError("未成功解析出任何 CRA 量表数据")

        df_out = pd.DataFrame(all_rows)
        df_out.to_excel(output_excel_path, index=False)
        log_fn(f"[成功] CRA 横表合并完成，共汇总 {len(all_rows)} 份数据 -> {output_excel_path}")


# ---------- 4. CRA 竖表合并算法 ----------

def execute_cra_vertical(archive_path, output_excel_path, log_fn=print):
    log_fn("[信息] 开始执行 CRA 竖表明细合并...")
    with tempfile.TemporaryDirectory() as tmpdir:
        files = extract_archive_to_dir(archive_path, tmpdir)
        data_files = [f for f in files if f.lower().endswith(('.xlsx', '.xlsm')) and not os.path.basename(f).startswith('~$')]
        
        detail_rows = []
        for f in data_files:
            base_name = os.path.basename(f)
            try:
                wb = openpyxl.load_workbook(f, data_only=True)
                if 'CRA填写表单' not in wb.sheetnames:
                    continue
                ws = wb['CRA填写表单']
                h1 = str(ws.cell(1, 2).value or '')
                h2 = str(ws.cell(1, 4).value or '')
                h3 = str(ws.cell(1, 6).value or '')

                for r in range(3, ws.max_row + 1):
                    times = to_num(ws.cell(r, 8).value)
                    if times is None or times == 0:
                        continue
                    row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 30))]
                    detail_rows.append([base_name, h1, h2, h3] + row_vals)
            except Exception as e:
                log_fn(f"[警告] 读取异常 {base_name}: {e}")

        if not detail_rows:
            raise ValueError("未提取到任何有效流水明细行")

        headers = ['来源文件', '第1行第2列', '第1行第4列', '第1行第6列'] + [f'列{i}' for i in range(1, len(detail_rows[0]) - 3)]
        df_res = pd.DataFrame(detail_rows, columns=headers[:len(detail_rows[0])])
        df_res.to_excel(output_excel_path, index=False)
        log_fn(f"[成功] CRA 竖表合并完成！共 {len(detail_rows)} 行明细 -> {output_excel_path}")


# ============================================================
#                      图形界面 (GUI)
# ============================================================

class DesktopApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Excel Matrix Pro - Windows 独立桌面版")
        self.geometry("860x620")
        self.minsize(760, 520)
        self.configure(bg="#f8fafc")
        self._setup_style()
        self._build_ui()

    def _setup_style(self):
        self.style = ttk.Style(self)
        self.style.theme_use('clam')
        self.style.configure('.', font=('Segoe UI', 9), background='#f8fafc')
        self.style.configure('TNotebook', background='#f8fafc')
        self.style.configure('TNotebook.Tab', padding=[14, 6], font=('Segoe UI', 9, 'bold'))
        self.style.configure('Accent.TButton', font=('Segoe UI', 9, 'bold'), foreground='#ffffff', background='#2563eb')
        self.style.map('Accent.TButton', background=[('active', '#1d4ed8'), ('disabled', '#93c5fd')])

    def _build_ui(self):
        # 顶部标题栏
        top_bar = tk.Frame(self, bg='#1e293b', height=45)
        top_bar.pack(fill=tk.X)
        lbl_title = tk.Label(top_bar, text="Excel Matrix Pro 批量处理平台 (免安装版)", font=('Segoe UI', 12, 'bold'), bg='#1e293b', fg='#ffffff')
        lbl_title.pack(side=tk.LEFT, padx=16, pady=10)

        # Notebook 选项卡
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=12, pady=10)

        # Tab 1: PM 拆分
        self.tab_split = ttk.Frame(self.notebook, padding=14)
        self.notebook.add(self.tab_split, text="PM 质量分拆分")
        self._init_split_tab()

        # Tab 2: PM 合并
        self.tab_merge = ttk.Frame(self.notebook, padding=14)
        self.notebook.add(self.tab_merge, text="PM 质量分合并")
        self._init_merge_tab()

        # Tab 3: CRA 横表
        self.tab_cra = ttk.Frame(self.notebook, padding=14)
        self.notebook.add(self.tab_cra, text="CRA 量表横表")
        self._init_cra_tab()

        # Tab 4: CRA 竖表
        self.tab_crav = ttk.Frame(self.notebook, padding=14)
        self.notebook.add(self.tab_crav, text="CRA 量表竖表")
        self._init_crav_tab()

        # 底部日志控制台
        console_frame = tk.LabelFrame(self, text="执行控制台日志", font=('Segoe UI', 9, 'bold'), bg='#f8fafc', fg='#334155', padx=8, pady=4)
        console_frame.pack(fill=tk.X, padx=12, pady=(0, 10))
        self.log_text = tk.Text(console_frame, height=7, font=('Consolas', 9), bg='#0f172a', fg='#f8fafc', insertbackground='white')
        self.log_text.pack(fill=tk.BOTH, expand=True)

    def log(self, msg):
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)

    def _file_row(self, parent, label_text, entry_var, is_dir=False):
        f = ttk.Frame(parent)
        f.pack(fill=tk.X, pady=4)
        ttk.Label(f, text=label_text, width=16, anchor=tk.W).pack(side=tk.LEFT)
        ttk.Entry(f, textvariable=entry_var).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=6)
        def browse():
            if is_dir:
                p = filedialog.askdirectory()
            else:
                p = filedialog.askopenfilename(filetypes=[("支持的文件", "*.xlsx *.xls *.xlsm *.zip *.7z"), ("所有文件", "*.*")])
            if p:
                entry_var.set(p)
        ttk.Button(f, text="浏览...", command=browse, width=8).pack(side=tk.RIGHT)

    def _init_split_tab(self):
        p = self.tab_split
        self.split_data_var = tk.StringVar()
        self.split_tpl_var = tk.StringVar()
        self._file_row(p, "数据源表格:", self.split_data_var)
        self._file_row(p, "样式模板文件:", self.split_tpl_var)

        btn = ttk.Button(p, text="▶ 执行 PM 拆分并打包 (ZIP)", style='Accent.TButton', command=self._start_split)
        btn.pack(fill=tk.X, pady=16)

    def _init_merge_tab(self):
        p = self.tab_merge
        self.merge_zip_var = tk.StringVar()
        self._file_row(p, "合并源文件/压缩包:", self.merge_zip_var)
        btn = ttk.Button(p, text="▶ 执行 PM 质量分合并 (XLSX)", style='Accent.TButton', command=self._start_merge)
        btn.pack(fill=tk.X, pady=16)

    def _init_cra_tab(self):
        p = self.tab_cra
        self.cra_zip_var = tk.StringVar()
        self.cra_tpl_var = tk.StringVar()
        self._file_row(p, "CRA表单文件/包:", self.cra_zip_var)
        self._file_row(p, "标准参考模板 (可选):", self.cra_tpl_var)
        btn = ttk.Button(p, text="▶ 执行 CRA 横表展开合并 (XLSX)", style='Accent.TButton', command=self._start_cra)
        btn.pack(fill=tk.X, pady=16)

    def _init_crav_tab(self):
        p = self.tab_crav
        self.crav_zip_var = tk.StringVar()
        self._file_row(p, "CRA表单文件/包:", self.crav_zip_var)
        btn = ttk.Button(p, text="▶ 执行 CRA 竖表明细过滤合并 (XLSX)", style='Accent.TButton', command=self._start_crav)
        btn.pack(fill=tk.X, pady=16)

    # ---------- 异步任务线程 ----------
    def _run_async(self, target_fn):
        t = threading.Thread(target=target_fn, daemon=True)
        t.start()

    def _start_split(self):
        data_p = self.split_data_var.get()
        tpl_p = self.split_tpl_var.get()
        if not data_p or not tpl_p:
            messagebox.showwarning("提示", "请选择数据源文件与模板文件！")
            return
        out_p = filedialog.asksaveasfilename(defaultextension=".zip", filetypes=[("ZIP 压缩包", "*.zip")])
        if not out_p:
            return
        def worker():
            try:
                execute_pm_split(data_p, tpl_p, out_p, log_fn=self.log)
                messagebox.showinfo("完成", f"PM 拆分成功完成！\n保存于: {out_p}")
            except Exception as e:
                self.log(f"[错误] {e}")
                messagebox.showerror("出错", str(e))
        self._run_async(worker)

    def _start_merge(self):
        p = self.merge_zip_var.get()
        if not p:
            messagebox.showwarning("提示", "请选择待合并文件！")
            return
        out_p = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel 表格", "*.xlsx")])
        if not out_p:
            return
        def worker():
            try:
                execute_pm_merge(p, out_p, log_fn=self.log)
                messagebox.showinfo("完成", f"合并成功！\n保存于: {out_p}")
            except Exception as e:
                self.log(f"[错误] {e}")
                messagebox.showerror("出错", str(e))
        self._run_async(worker)

    def _start_cra(self):
        p = self.cra_zip_var.get()
        tpl = self.cra_tpl_var.get() or None
        if not p:
            messagebox.showwarning("提示", "请选择 CRA 表单文件！")
            return
        out_p = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel 表格", "*.xlsx")])
        if not out_p:
            return
        def worker():
            try:
                execute_cra_horizontal(p, out_p, template_path=tpl, log_fn=self.log)
                messagebox.showinfo("完成", f"CRA 横表合并成功！\n保存于: {out_p}")
            except Exception as e:
                self.log(f"[错误] {e}")
                messagebox.showerror("出错", str(e))
        self._run_async(worker)

    def _start_crav(self):
        p = self.crav_zip_var.get()
        if not p:
            messagebox.showwarning("提示", "请选择 CRA 表单文件！")
            return
        out_p = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel 表格", "*.xlsx")])
        if not out_p:
            return
        def worker():
            try:
                execute_cra_vertical(p, out_p, log_fn=self.log)
                messagebox.showinfo("完成", f"CRA 竖表合并成功！\n保存于: {out_p}")
            except Exception as e:
                self.log(f"[错误] {e}")
                messagebox.showerror("出错", str(e))
        self._run_async(worker)


if __name__ == '__main__':
    app = DesktopApp()
    app.mainloop()
